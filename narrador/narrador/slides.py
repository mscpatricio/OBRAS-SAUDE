"""Leitura de apresentacoes: imagens dos slides + roteiro de narracao.

Fontes aceitas:
  * .pptx  -> imagens via LibreOffice (pptx -> pdf -> png) e notas do orador
  * .pdf   -> imagens via PyMuPDF (roteiro vem de arquivo .md/.json/.yaml)
  * pasta  -> imagens .png/.jpg ordenadas por nome (roteiro em arquivo)

Formato do roteiro em Markdown (um bloco por slide):

    ## Slide 1
    Texto que sera narrado no slide 1.

    ## Slide 2
    Texto do slide 2. Paragrafos separados geram pausas maiores.

Tambem aceita JSON/YAML: lista de strings ou lista de {"slide": n, "texto": "..."}.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import yaml

from .util import ErroNarrador, console, exigir_binario, garantir_dir

IMAGENS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


@dataclass
class Slide:
    numero: int  # 1-based
    imagem: Path
    texto: str = ""
    titulo: str = ""


@dataclass
class Apresentacao:
    nome: str
    slides: list[Slide] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Imagens
# --------------------------------------------------------------------------- #

def _ordenar_natural(caminhos: list[Path]) -> list[Path]:
    def chave(p: Path):
        return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", p.stem)]
    return sorted(caminhos, key=chave)


def pptx_para_pdf(pptx: Path, destino_dir: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ErroNarrador(
            "LibreOffice nao encontrado. Instale (sudo apt install libreoffice-impress) "
            "ou exporte a apresentacao para PDF e passe o .pdf junto com um roteiro."
        )
    garantir_dir(destino_dir)
    perfil = tempfile.mkdtemp(prefix="lo_profile_")
    env = dict(os.environ, HOME=perfil)
    cmd = [
        soffice,
        f"-env:UserInstallation=file://{perfil}",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(destino_dir),
        str(pptx),
    ]
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, env=env, timeout=600)
    pdf = destino_dir / (pptx.stem + ".pdf")
    if proc.returncode != 0 or not pdf.exists():
        raise ErroNarrador(f"Falha ao converter PPTX em PDF com LibreOffice:\n{proc.stderr[-2000:]}")
    shutil.rmtree(perfil, ignore_errors=True)
    return pdf


def pdf_para_imagens(pdf: Path, destino_dir: Path, largura: int = 1920) -> list[Path]:
    import pymupdf  # import tardio: dependencia pesada

    garantir_dir(destino_dir)
    saidas: list[Path] = []
    with pymupdf.open(pdf) as doc:
        for i, pagina in enumerate(doc, start=1):
            escala = largura / pagina.rect.width
            pix = pagina.get_pixmap(matrix=pymupdf.Matrix(escala, escala), alpha=False)
            saida = destino_dir / f"slide_{i:03d}.png"
            pix.save(saida)
            saidas.append(saida)
    if not saidas:
        raise ErroNarrador(f"O PDF {pdf} nao tem paginas.")
    return saidas


def imagens_da_pasta(pasta: Path) -> list[Path]:
    arquivos = [p for p in pasta.iterdir() if p.suffix.lower() in IMAGENS]
    if not arquivos:
        raise ErroNarrador(f"Nenhuma imagem (.png/.jpg) encontrada em {pasta}.")
    return _ordenar_natural(arquivos)


# --------------------------------------------------------------------------- #
# Notas do orador e texto dos slides (PPTX)
# --------------------------------------------------------------------------- #

def notas_pptx(pptx: Path) -> list[tuple[str, str, str]]:
    """Retorna, por slide, (titulo, notas_do_orador, texto_visivel)."""
    from pptx import Presentation

    prs = Presentation(str(pptx))
    resultado = []
    for slide in prs.slides:
        titulo = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            titulo = slide.shapes.title.text_frame.text.strip()
        notas = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notas = slide.notes_slide.notes_text_frame.text.strip()
        textos = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                t = shape.text_frame.text.strip()
                if t:
                    textos.append(t)
        resultado.append((titulo, notas, "\n".join(textos)))
    return resultado


# --------------------------------------------------------------------------- #
# Roteiro externo
# --------------------------------------------------------------------------- #

_RE_CABECALHO = re.compile(r"^\s*#{1,3}\s*(?:slide|slaide|p[aá]gina|cena)?\s*(\d+)\b.*$", re.IGNORECASE | re.MULTILINE)


def ler_roteiro(caminho: Path) -> dict[int, str]:
    """Le roteiro .md/.txt/.json/.yaml e devolve {numero_slide: texto}."""
    if not caminho.exists():
        raise ErroNarrador(f"Roteiro nao encontrado: {caminho}")
    bruto = caminho.read_text(encoding="utf-8")
    suf = caminho.suffix.lower()
    if suf in {".json", ".yaml", ".yml"}:
        dados = json.loads(bruto) if suf == ".json" else yaml.safe_load(bruto)
        if isinstance(dados, dict) and "slides" in dados:
            dados = dados["slides"]
        mapa: dict[int, str] = {}
        for i, item in enumerate(dados or [], start=1):
            if isinstance(item, str):
                mapa[i] = item.strip()
            elif isinstance(item, dict):
                n = int(item.get("slide", item.get("numero", i)))
                mapa[n] = str(item.get("texto", item.get("text", ""))).strip()
        return mapa

    # Markdown / texto: cabecalhos "## Slide N"; sem cabecalhos, separa por "---" ou linhas em branco duplas
    cabecalhos = list(_RE_CABECALHO.finditer(bruto))
    mapa = {}
    if cabecalhos:
        for idx, m in enumerate(cabecalhos):
            fim = cabecalhos[idx + 1].start() if idx + 1 < len(cabecalhos) else len(bruto)
            mapa[int(m.group(1))] = bruto[m.end():fim].strip()
        return mapa
    blocos = [b.strip() for b in re.split(r"^\s*---+\s*$", bruto, flags=re.MULTILINE) if b.strip()]
    if len(blocos) == 1:
        blocos = [b.strip() for b in re.split(r"\n\s*\n\s*\n", bruto) if b.strip()]
    return {i: b for i, b in enumerate(blocos, start=1)}


def escrever_roteiro_md(apr: Apresentacao, destino: Path) -> Path:
    linhas = [f"# Roteiro de narracao: {apr.nome}", "",
              "Edite o texto abaixo de cada slide. Paragrafos separados geram pausas maiores.", ""]
    for s in apr.slides:
        cab = f"## Slide {s.numero}" + (f" - {s.titulo}" if s.titulo else "")
        linhas += [cab, "", s.texto or "(sem texto: escreva a narracao aqui)", ""]
    destino.write_text("\n".join(linhas), encoding="utf-8")
    return destino


# --------------------------------------------------------------------------- #
# Montagem
# --------------------------------------------------------------------------- #

def carregar(origem: Path, trabalho: Path, roteiro: Path | None = None,
             usar_texto_slide: bool = True, largura: int = 1920) -> Apresentacao:
    """Carrega imagens + textos de narracao a partir da origem.

    Prioridade do texto: roteiro externo > notas do orador > texto visivel do slide.
    """
    origem = origem.expanduser().resolve()
    if not origem.exists():
        raise ErroNarrador(f"Arquivo ou pasta nao encontrado: {origem}")
    dir_slides = garantir_dir(trabalho / "slides")
    nome = origem.stem if origem.is_file() else origem.name
    apr = Apresentacao(nome=nome)

    metadados: list[tuple[str, str, str]] = []
    if origem.is_dir():
        imagens = imagens_da_pasta(origem)
    elif origem.suffix.lower() == ".pptx":
        console.print("[cyan]Convertendo PPTX em imagens com LibreOffice...[/]")
        pdf = pptx_para_pdf(origem, trabalho)
        imagens = pdf_para_imagens(pdf, dir_slides, largura)
        metadados = notas_pptx(origem)
    elif origem.suffix.lower() == ".pdf":
        imagens = pdf_para_imagens(origem, dir_slides, largura)
    elif origem.suffix.lower() in IMAGENS:
        imagens = [origem]
    else:
        raise ErroNarrador(f"Formato nao suportado: {origem.suffix}. Use .pptx, .pdf ou pasta de imagens.")

    textos_roteiro = ler_roteiro(roteiro) if roteiro else {}

    for i, img in enumerate(imagens, start=1):
        titulo, notas, visivel = metadados[i - 1] if i - 1 < len(metadados) else ("", "", "")
        texto = textos_roteiro.get(i) or notas or (visivel if usar_texto_slide else "")
        apr.slides.append(Slide(numero=i, imagem=img, texto=texto.strip(), titulo=titulo))

    if metadados and len(metadados) != len(imagens):
        console.print(f"[yellow]Aviso:[/] {len(metadados)} slides no PPTX, {len(imagens)} paginas no PDF. "
                      "Slides ocultos podem causar desalinhamento; confira o roteiro gerado.")
    sem_texto = [s.numero for s in apr.slides if not s.texto]
    if sem_texto:
        console.print(f"[yellow]Aviso:[/] slides sem narracao (ficarao em silencio): {sem_texto}")
    return apr
