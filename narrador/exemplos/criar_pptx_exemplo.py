"""Gera uma apresentacao de exemplo (3 slides com notas do orador) para testar o Narrador.

    python exemplos/criar_pptx_exemplo.py saida/exemplo.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDES = [
    ("Obras de Saúde - SESA", "Acompanhamento das obras em 2026",
     "Olá! Nesta apresentação vamos acompanhar a execução das obras de saúde da SESA no exercício de 2026."),
    ("Panorama geral", "12 obras em andamento\nR$ 48.350.000,00 contratados\n63,5% executado até 30/09/2026",
     "Hoje são 12 obras em andamento, com investimento total de R$ 48.350.000,00.\n\nAté 30/09/2026, 63,5% do valor contratado já foi executado."),
    ("Próximos passos", "Medições em 15/11/2026\nPrestação de contas",
     "As próximas medições acontecem em 15/11/2026. Fiquem atentos aos prazos de prestação de contas.\n\nObrigado pela atenção!"),
]


def criar(destino: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for titulo, corpo, notas in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = titulo
        caixa = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11), Inches(4))
        tf = caixa.text_frame
        tf.word_wrap = True
        for i, linha in enumerate(corpo.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = linha
            p.font.size = Pt(32)
            p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        slide.notes_slide.notes_text_frame.text = notas
    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destino))
    return destino


if __name__ == "__main__":
    alvo = Path(sys.argv[1] if len(sys.argv) > 1 else "saida/exemplo.pptx")
    print(criar(alvo))
